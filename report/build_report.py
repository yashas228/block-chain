"""Build the final project report DOCX.

Run:  python report/build_report.py
Output: SP_FINAL_REPORT.docx (repository root)
"""
import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
import docx_kit as K  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(ROOT, "report", "figures")
OUT = os.path.join(ROOT, "SP_FINAL_REPORT.docx")

TITLE = "Blockchain-Based Academic Record Management and Verification System"
SUBTITLE = "A Web3 Student Information System with Zero-Gas Credential Verification"

STUDENTS = [
    ("Kasala Yashas Reddy", "24STUCHH010288"),
    ("Kurri Bhavesh Kumar", "24STUCHH010966"),
    ("Vegesna Sujit Varma", "24STUCHH010891"),
    ("Mantena Laxmi Kishore Varma", "24STUCHH010060"),
]


def figure_path(name):
    return os.path.join(FIG, name)


# ------------------------------------------------------------------ front matter
def title_page(doc):
    K.blank(doc, 1)
    K.centered(doc, "THE ICFAI FOUNDATION FOR HIGHER EDUCATION, HYDERABAD",
               size=13, bold=True, color=K.NAVY, space_after=2)
    K.centered(doc, "FACULTY OF SCIENCE AND TECHNOLOGY", size=11, bold=True,
               color=K.STEEL, space_after=18)
    K.centered(doc, "A FINAL PROJECT REPORT ON", size=10.5, color=K.GREY, space_after=14)
    K.centered(doc, TITLE, size=20, bold=True, color=K.NAVY, space_after=6, font="Cambria")
    K.centered(doc, SUBTITLE, size=12, italic=True, color=K.STEEL, space_after=22)
    K.centered(doc, "Submitted in partial fulfilment of the requirements for the award of the degree of",
               size=10.5, space_after=4)
    K.centered(doc, "BACHELOR OF TECHNOLOGY", size=14, bold=True, color=K.NAVY, space_after=2)
    K.centered(doc, "SPECIAL PROJECT (SP301)", size=11, bold=True, color=K.STEEL, space_after=22)
    K.centered(doc, "Submitted by", size=10.5, italic=True, space_after=8)

    for name, roll in STUDENTS:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.1
        run = para.add_run(name)
        run.font.size = Pt(12)
        run.bold = True
        run.font.name = K.BODY_FONT
        run = para.add_run(f"   ({roll})")
        run.font.size = Pt(11)
        run.font.name = K.BODY_FONT

    K.blank(doc, 1)
    K.centered(doc, "Under the guidance of", size=10.5, italic=True, space_after=4)
    K.centered(doc, "Dr. D. Krishna Madhuri", size=13, bold=True, color=K.NAVY, space_after=18)
    K.centered(doc, "Department of Computer Science and Engineering", size=10.5, space_after=2)
    K.centered(doc, "October 2026", size=11, bold=True, color=K.STEEL, space_after=0)


def declaration(doc, page_break=True):
    K.front_heading(doc, "Declaration")
    K.p(doc, "We hereby declare that the project report entitled "
             f"*“{TITLE}”* submitted to the Faculty of Science and Technology, "
             "ICFAI Foundation for Higher Education, Hyderabad, in partial fulfilment of the "
             "requirements for the award of the degree of Bachelor of Technology, is a record of "
             "original work carried out by us under the guidance of Dr. D. Krishna Madhuri.")
    K.p(doc, "The system described in this report was designed, implemented and evaluated by the "
             "authors. All external sources of information, standards, libraries and research "
             "publications that informed the design are acknowledged in the references. The "
             "measurements reported in Chapter 6 were produced by the measurement scripts contained "
             "in the project repository, and the raw values are retained in the repository so that "
             "the results can be reproduced.")
    K.p(doc, "This work has not been submitted, either in part or in full, for the award of any "
             "other degree or diploma at this or any other institution.")
    K.blank(doc, 2)

    from docx.shared import Inches
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = True
    cells = tbl.rows[0].cells
    for cell, (name, roll) in zip(cells, STUDENTS[:2]):
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(18)
        run = para.add_run("Signature: ______________________\n")
        run.font.size = Pt(10)
        run = para.add_run(f"{name}\n{roll}")
        run.font.size = Pt(10)
        run.bold = True
    K.blank(doc, 1)
    tbl = doc.add_table(rows=1, cols=2)
    cells = tbl.rows[0].cells
    for cell, (name, roll) in zip(cells, STUDENTS[2:]):
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(18)
        run = para.add_run("Signature: ______________________\n")
        run.font.size = Pt(10)
        run = para.add_run(f"{name}\n{roll}")
        run.font.size = Pt(10)
        run.bold = True
    K.blank(doc, 2)
    K.p(doc, "Place: Hyderabad", size=10.5, space_after=2, align="left")
    K.p(doc, "Date:", size=10.5, align="left")


def certificate(doc):
    K.front_heading(doc, "Certificate")
    K.p(doc, "This is to certify that the project report entitled "
             f"*“{TITLE}”* is a bonafide record of the special project work carried out by "
             "Kasala Yashas Reddy (24STUCHH010288), Kurri Bhavesh Kumar (24STUCHH010966), "
             "Vegesna Sujit Varma (24STUCHH010891) and Mantena Laxmi Kishore Varma "
             "(24STUCHH010060), students of the Faculty of Science and Technology, ICFAI Foundation "
             "for Higher Education, Hyderabad, during the academic year 2026–27, under my "
             "supervision and guidance.")
    K.p(doc, "The work presented in this report is original and has been evaluated through the test "
             "suites and measurement scripts described in Chapters 5 and 6. It has not been "
             "submitted elsewhere for the award of any degree or diploma.")
    K.blank(doc, 3)
    K.p(doc, "Project Guide: ______________________", size=10.5, align="left", space_after=2)
    K.p(doc, "Dr. D. Krishna Madhuri", size=10, align="left", space_after=18)
    K.p(doc, "Head of Department: ______________________", size=10.5, align="left", space_after=2)
    K.p(doc, "Department of Computer Science and Engineering", size=10, align="left")


def acknowledgement(doc):
    K.front_heading(doc, "Acknowledgement")
    K.p(doc, "We express our sincere gratitude to our project guide, Dr. D. Krishna Madhuri, for "
             "her continuous guidance, careful review of our design decisions and her insistence "
             "that every claim in this report be backed by a measurement or a test result. Her "
             "feedback shaped both the architecture and the way we evaluated it.")
    K.p(doc, "We thank the faculty and staff of the Department of Computer Science and Engineering, "
             "Faculty of Science and Technology, ICFAI Foundation for Higher Education, Hyderabad, "
             "for the laboratory facilities and the academic environment that made this work "
             "possible. We are grateful to our classmates who acted as test users for the issuer "
             "and verification screens, and who reported several usability problems during the "
             "demonstrations.")
    K.p(doc, "Finally, we thank our families for their patience and support during the many long "
             "evenings that this project required. Any errors that remain in the implementation or "
             "in this report are ours alone.")


def abstract(doc):
    K.front_heading(doc, "Abstract")
    K.p(doc, "Academic credentials are still largely issued and verified through paper documents and "
             "centralised institutional databases. Verification therefore depends on direct contact "
             "between the verifier and the issuing institution, is slow, and gives a verifier no "
             "independent way to tell whether a presented transcript has been altered. This report "
             "presents the design, implementation and evaluation of a blockchain-based academic "
             "record management and verification system that keeps the institution's operational "
             "data off-chain and anchors only cryptographic evidence on an Ethereum-compatible "
             "ledger.")
    K.p(doc, "The implemented system is a full-stack student information system. Academic data — "
             "students, programmes, courses, enrolments and grades — live in a MySQL 8 database, "
             "which remains the mutable source of truth. When a registrar issues a transcript, the "
             "backend compiles a canonical, key-sorted JSON document from the live grade data, "
             "computes its SHA-256 digest (or the digest of an uploaded PDF), encrypts the stored "
             "document with AES-256-GCM, and anchors the digest, storage reference and issuer "
             "identity in an AcademicRecordManager smart contract. A record is marked ACTIVE only "
             "after the anchoring transaction is mined, so a failed transaction can never be "
             "presented as an issued credential. Corrections publish the corrected record and retire "
             "its predecessor in a single atomic transaction, and revocations update the on-chain "
             "status while preserving the audit trail.")
    K.p(doc, "Verification is deliberately treated as an integrated, zero-cost feature rather than a "
             "separate product. A verifier supplies a record identifier, a shared link or the "
             "document itself; the system reads the anchored evidence with a read-only `eth_call` "
             "that consumes no gas and requires no wallet, cross-checks the ledger against the "
             "database, and reports the database status, the ledger status, an explicit consistency "
             "flag and the issuer's signature check. Unknown identifiers return 404 and an "
             "unreachable chain is reported as unavailable, never as tampering. Identity is "
             "wallet-based: users sign a login challenge with MetaMask (or another EVM wallet) and "
             "the API issues a stateless twelve-hour session token, while the contract independently "
             "enforces its own issuer whitelist.")
    K.p(doc, "The system was evaluated with fourteen contract unit tests, a fifty-two assertion "
             "end-to-end API suite, a ten assertion chain-service round trip, a twenty-three "
             "mutation tamper suite, and a seven assertion chain-outage test, all of which passed. "
             "Public verification required 45–69 ms (mean 54.7 ms, n = 10) at zero gas; issuance "
             "required a mean of 304 ms and 101,548 gas; a revocation cost 53,179 gas and a "
             "correction 150,684 gas; a pre-signed batch measured 194.2 transactions per second "
             "against a 20 tx/s acceptance target. The report also states the limits of the "
             "prototype: the deployment is a local Hardhat network, Sepolia and IPFS pinning are "
             "configured but were not exercised, the credential signature is a wallet signature over "
             "the record digest rather than a W3C Data Integrity proof, and the trust model still "
             "assumes a single administrative authority.")


def contents(doc):
    doc.add_page_break()
    K.toc(doc, "Table of Contents", 'TOC \\o "1-3" \\h \\z \\u',
          "Open in Word and update this field to generate the table of contents with page numbers.")
    doc.add_page_break()
    K.toc(doc, "List of Figures", 'TOC \\h \\z \\t "FigCaption,1"',
          "Open in Word and update this field to list the figures with page numbers.")
    doc.add_page_break()
    K.toc(doc, "List of Tables", 'TOC \\h \\z \\t "TabCaption,1"',
          "Open in Word and update this field to list the tables with page numbers.")


# ---------------------------------------------------------------------- chapter 1
def chapter1(doc):
    K.h1(doc, "1. Introduction")

    K.h2(doc, "1.1 Background and motivation")
    K.p(doc, "A degree certificate, a transcript or a mark sheet is the interface between years of "
             "study and the opportunities that follow it. Employers, admissions committees, "
             "scholarship bodies, immigration authorities and professional councils all need to "
             "answer the same question: is this document genuine, and does it still represent the "
             "institution's position? In most institutions that question is answered by a "
             "centralised database that only the institution can query, or by a paper document "
             "carrying a seal and a signature that must be verified manually.")
    K.p(doc, "That arrangement has three structural weaknesses. First, verification is a service "
             "the issuer must perform: an employer who receives a transcript from a candidate has to "
             "contact the university, wait for a response, and trust the channel through which the "
             "reply arrives. Second, the artefact and the record can diverge. Once a transcript "
             "leaves the registrar's office it is a file, and a file can be edited; nothing on the "
             "paper or in the PDF lets a third party detect that the grade point average printed on "
             "it was changed after issuance. Third, the lifecycle of a credential is invisible. A "
             "degree that has been revoked after a disciplinary finding, or a transcript that has "
             "been superseded after a correction, looks identical to a valid one when it is "
             "presented months later.")
    K.p(doc, "Public blockchains offer a property that centralised databases do not: an append-only "
             "ledger whose contents any participant can independently re-derive and check. If a "
             "digest of a credential is written to such a ledger at the moment of issuance, then a "
             "verifier who recomputes the digest of the presented document can determine whether the "
             "bytes are the bytes the institution anchored, without contacting the institution and "
             "without trusting the intermediary that delivered the file. The ledger also provides a "
             "natural place to record status changes, so revocation becomes a public, timestamped "
             "fact rather than an internal database row.")
    K.p(doc, "Blockchain does not, however, make a credential true, and it is important to be "
             "precise about what it does provide. It cannot validate the grade that a faculty member "
             "entered, it cannot make personal data disappear once written, and it cannot by itself "
             "give a verifier a machine-readable credential that interoperates with another "
             "institution's system. The literature review in Chapter 2 shows that the systems which "
             "have survived contact with real institutions treat the ledger as one component of a "
             "larger credential architecture, alongside digital signatures, identity, off-chain "
             "storage, status management and standardised data models. This project adopts that "
             "position: the blockchain is the anchoring and audit layer of a conventional student "
             "information system, not a replacement for it.")

    K.h2(doc, "1.2 Problem definition")
    K.p(doc, "The project addresses the following concrete problems in academic record management:")
    K.bullets(doc, [
        "**Certificate forgery.** Academic documents circulate as PDFs and scanned images that can "
        "be edited with ordinary tools. A verifier holding only the document has no independent "
        "means of establishing that its contents match what the institution issued.",
        "**Slow, issuer-dependent verification.** Confirming a credential usually requires a "
        "request to the registrar, manual lookup and a reply. The delay is measured in days and the "
        "process does not scale during admission or recruitment cycles.",
        "**Undetectable post-issuance modification.** In a centralised store, a change to a grade or "
        "a status leaves no externally checkable evidence. Even where audit tables exist, they are "
        "visible only to the institution.",
        "**Incomplete credential lifecycle.** Revocation, correction and supersession are poorly "
        "represented. A verifier cannot easily distinguish a credential that was never valid from "
        "one that was valid and has since been revoked.",
        "**Single point of failure and availability.** Verification stops when the institution's "
        "system is unavailable, and long-term availability depends on the institution continuing to "
        "operate the same database.",
        "**Weak student control.** Graduates depend on the institution to send transcripts on their "
        "behalf and have no portable, self-serviceable proof of their own academic record.",
    ])

    K.h2(doc, "1.3 Objectives")
    K.p(doc, "The objectives of the project, restated here so that Chapter 6 can trace results back "
             "to them, are:")
    K.bullets(doc, [
        "**O1 — Manage academic records end to end.** Implement issuance, retrieval, correction, "
        "revocation and transcript compilation as first-class operations of a student information "
        "system, driven by real programme, course, enrolment and grade data.",
        "**O2 — Make unauthorised modification detectable.** Anchor a cryptographic digest of the "
        "exact issued bytes on-chain and recompute it during verification, so that any change to the "
        "document is detectable by a third party.",
        "**O3 — Enable verification without the issuer and without cost.** Provide a public, "
        "wallet-free verification path that answers in well under a second and consumes no gas.",
        "**O4 — Preserve the credential lifecycle.** Represent Active, Revoked and Superseded states "
        "on-chain, link a superseded record to its replacement, and keep the transition atomic.",
        "**O5 — Give students portable access.** Let a student view and download their own "
        "credentials, and share a link or QR code that a verifier can resolve directly.",
        "**O6 — Enforce authorisation independently of the API.** Use role-based access control in "
        "the application and an issuer whitelist in the contract, so that a compromised or bypassed "
        "API still cannot write to the ledger.",
        "**O7 — Fail honestly.** Distinguish “not found”, “tampered”, “revoked” and “service "
        "unavailable” in both the API and the user interface, and never report an outage as "
        "tampering or a failed transaction as an issued credential.",
        "**O8 — Produce reproducible evidence.** Automate the functional, adversarial and performance "
        "evaluation so that every number quoted in this report can be regenerated from the "
        "repository.",
    ])

    K.h2(doc, "1.4 Scope of the delivered system")
    K.p(doc, "The scope was fixed after the midterm review and is summarised below. Items in the "
             "right-hand column are deliberately excluded and are revisited in Section 7.4.")
    K.table(doc,
            ["In scope (implemented and exercised)", "Out of scope (not implemented)"],
            [["Wallet-based sign-in and role-based access for ADMIN, ISSUER, STUDENT and VERIFIER",
              "Federated identity, institutional SSO and DID-based issuer identity"],
             ["Programme, course, enrolment and grade management with weighted GPA computation",
              "Integration with commercial student information systems or LMS platforms"],
             ["Credential issuance from live academic data or an uploaded PDF, with canonical JSON",
              "W3C Verifiable Credential issuance with Data Integrity proofs"],
             ["SHA-256 digest anchoring and AES-256-GCM encryption of stored documents",
              "On-chain storage of personal data or full credential documents"],
             ["Active / Revoked / Superseded lifecycle with atomic supersession",
              "Multi-institution consortium governance and cross-institution credit transfer"],
             ["Zero-gas public verification by identifier or by document hash",
              "Zero-knowledge selective disclosure of individual claims"],
             ["QR codes and shareable verification links",
              "Camera-based QR scanning and a native mobile application"],
             ["Local Hardhat deployment with Sepolia configuration present",
              "A funded public testnet or mainnet deployment with measured public-network cost"]],
            "Scope of the implemented system.",
            widths=[3.3, 3.2])

    K.h2(doc, "1.5 Contributions of the work")
    K.bullets(doc, [
        "A working hybrid architecture in which academic data stays in a relational database and the "
        "ledger stores only a digest, a storage reference, an issuer address and a status, with a "
        "documented rule for exactly which bytes are hashed.",
        "An issuance protocol that couples the database and the ledger correctly: a record is "
        "created as DRAFT, anchored, and only promoted to ACTIVE after the receipt is confirmed, "
        "with an explicit retry path for failures.",
        "A verification procedure that cross-checks two independent sources and reports a "
        "consistency flag, so that tampering with the database behind the ledger's back is visible "
        "rather than silently trusted.",
        "An adversarial evaluation, including a twenty-three mutation tamper suite and an automated "
        "chain-outage test, and a measured performance characterisation of the anchoring path.",
        "A usable interface for registrars and students, including QR-based sharing, with the "
        "verification page accessible without a wallet or an account.",
    ])

    K.h2(doc, "1.6 Limitations of the delivered system")
    K.p(doc, "The prototype is honest about what it does not do. These limitations are developed "
             "further in Section 7.3.")
    K.bullets(doc, [
        "All measurements were taken on a local Hardhat node with automining; they characterise the "
        "implementation, not a public network, where block times and fees dominate.",
        "Sepolia deployment and Pinata pinning are configured in the codebase but were not exercised "
        "because no funded testnet key or pinning credential was available during the project.",
        "The issuer credential signature is an EIP-191 wallet signature over the record digest. It "
        "is not a W3C Data Integrity proof, and the credential JSON is W3C-VC-shaped rather than "
        "conformant.",
        "The trust model is single-institution: the contract administrator, which is also the server "
        "signer in the default configuration, controls the issuer whitelist.",
        "QR codes are generated and resolve to the verification page, but camera-based scanning is "
        "not implemented; a verifier types or follows the link.",
        "The system anchors what the institution asserts. It cannot detect an incorrect grade "
        "entered before issuance, which is a governance problem rather than a cryptographic one.",
    ])

    K.h2(doc, "1.7 Organisation of the report")
    K.p(doc, "Chapter 2 reviews the research literature and the standards landscape, and derives the "
             "design requirements used later. Chapter 3 analyses the existing process, states the "
             "requirements and records the feasibility and environment decisions. Chapter 4 presents "
             "the architecture, data model, data-flow and UML design. Chapter 5 describes the "
             "implementation of the contract, the API, the storage layer and the client. Chapter 6 "
             "reports the testing strategy and the measured results, including the threats to their "
             "validity. Chapter 7 concludes and sets out the work that should follow. Appendices "
             "list the REST interface, the commands needed to reproduce the evidence, and the "
             "abbreviations used throughout.")


# ---------------------------------------------------------------------- chapter 2
def chapter2(doc):
    K.h1(doc, "2. Literature Review")

    K.h2(doc, "2.1 Review approach")
    K.p(doc, "The review was conducted to answer a narrow question: which architectural choices in "
             "a blockchain-based academic credential system are supported by evidence, and which are "
             "assertions that the literature does not support? Peer-reviewed studies, standards "
             "documents and the documentation of deployed open-source systems were considered. "
             "Marketing material and unverifiable claims were excluded. The anchor study is the "
             "systematic review by Rustemi et al., which applied a PRISMA protocol to 1,744 "
             "publications between 2018 and 2022 and retained 34 primary studies on blockchain-based "
             "academic certificate verification [1]. That review is useful precisely because it is "
             "deflationary: it finds evidence of benefit in reliability, efficiency and independence "
             "from a central authority, but concludes that the field is not mature and that "
             "non-technical factors dominate adoption.")
    K.p(doc, "Three standards documents frame the analysis: the W3C Verifiable Credentials Data "
             "Model 2.0, published as a Recommendation in May 2025 [2]; the Open Badges 3.0 "
             "specification from 1EdTech, which profiles verifiable credentials for achievement "
             "claims [12]; and the European Commission's report on blockchain in education, which "
             "sets out the policy and ecosystem view [3][9]. The distinction the literature insists "
             "on, and which this project adopts, is that a blockchain is an anchoring technology, "
             "while a credential format is a data and cryptographic model. A system needs both, and "
             "one does not substitute for the other.")

    K.h2(doc, "2.2 How academic credentials are verified today")
    K.p(doc, "Conventional verification relies on a database lookup performed by the issuing "
             "institution, on paper artefacts carrying seals and signatures, or on e-mail exchanges "
             "between registrars. Each of these depends on the issuing institution remaining "
             "reachable and cooperative, and each places the verifier in the position of trusting "
             "whatever channel delivered the response. The cost is not only delay: because the "
             "verifier cannot check anything independently, verification effort is concentrated at "
             "the issuer, which is exactly where capacity is scarcest during admission cycles. The "
             "systematic review identifies this dependency as the central inefficiency that "
             "blockchain-based proposals attempt to remove [1][3].")
    K.p(doc, "A second, subtler problem is that verification answers the wrong question. A "
             "registrar's confirmation typically establishes that a credential was issued, not that "
             "the document in the verifier's hands is the document that was issued. Where systems do "
             "bind the artefact to the record, they do so with a digital signature whose verification "
             "requires the issuer's public key and a signature format the verifier must implement. "
             "Blockchain anchoring simplifies the trust distribution problem — the digest is "
             "published, so anyone who can hash a file can check it — but it does not remove the need "
             "for signatures or for a defined serialisation of the signed bytes.")

    K.h2(doc, "2.3 Thematic findings")

    K.h3(doc, "2.3.1 Immutability does not establish truth")
    K.p(doc, "The most consistently repeated caution in the literature is that a ledger records "
             "assertions, not facts. If an authorised issuer anchors an incorrect grade, the ledger "
             "preserves the error with perfect fidelity. Blockcerts documentation makes the same "
             "point from the opposite direction: a verifiable credential proves that a claim was "
             "made by an issuer and has not been altered since, not that the claim is correct [4][5]. "
             "Any design must therefore treat issuer authorisation and data-entry governance as "
             "part of the security boundary. In this project the boundary is explicit: the digest "
             "proves integrity after issuance, and the database plus institutional process remain "
             "responsible for correctness at issuance.")

    K.h3(doc, "2.3.2 Personal data should not be written to a public ledger")
    K.p(doc, "Academic records contain names, identifiers, grades and sometimes sensitive personal "
             "circumstances. Because public ledger entries are replicated and effectively permanent, "
             "writing such data on-chain conflicts with data-protection expectations and with the "
             "right to erasure. The design consensus is to store the credential off-chain and anchor "
             "a commitment — a hash, an identifier, a status — on-chain. ZKBAR-V extends this idea "
             "with a dual-blockchain design, IPFS storage, decentralised identifiers and "
             "zero-knowledge proofs, so that a verifier can establish a property of a record without "
             "receiving the record [6]. The cost is computational complexity and a considerably "
             "larger trust surface, which is why this project stops at the commitment and leaves "
             "zero-knowledge disclosure to future work.")

    K.h3(doc, "2.3.3 The credential lifecycle is the hard part")
    K.p(doc, "Issuance is the easy case. A production credential system must also answer whether a "
             "credential is currently valid, must support correction when a grade is revised, and "
             "must survive key rotation. W3C Verifiable Credentials 2.0 devotes explicit machinery to "
             "credential status, and Blockcerts defines revocation and status handling for its "
             "certificates [2][5]. The lesson for this project is that status must be a first-class "
             "on-chain field rather than an inference from the presence of a hash, and that a "
             "correction should not merely delete the old record: it must leave a link from the "
             "superseded credential to its replacement, so that a verifier holding an older copy can "
             "discover what happened to it.")

    K.h3(doc, "2.3.4 Interoperability requires standards, not clever schemas")
    K.p(doc, "Early blockchain credential projects defined their own JSON schemas and their own "
             "signature formats, which made them difficult to consume outside the issuing platform. "
             "The current consensus is towards W3C Verifiable Credentials, Decentralised "
             "Identifiers, the Data Integrity proof suites and Open Badges 3.0 [2][12]. EduCTX shows "
             "the same issue at the level of academic data: its contribution is a shared "
             "representation of ECTS credits across institutions, which only has value if "
             "institutions agree on the representation [7]. Interoperability is therefore a "
             "governance problem supported by standards, not a feature that a single implementation "
             "can deliver on its own.")

    K.h3(doc, "2.3.5 Scale, cost and governance")
    K.p(doc, "Public networks impose fees, variable confirmation latency and external capacity "
             "constraints. Permissioned and consortium networks offer predictable throughput and "
             "institutional control, at the cost of membership and arbitration agreements that the "
             "literature identifies as unresolved [1][7]. There is no evidence in the reviewed work "
             "for a universally best consensus mechanism for academic credentials. The practical "
             "consequence for a prototype is to keep the on-chain footprint minimal — one write per "
             "issuance, one write per lifecycle change, and reads that cost nothing — so that the "
             "design remains viable whether it is later deployed to a Layer 2 network or to a "
             "consortium chain.")

    K.h3(doc, "2.3.6 Usability and adoption")
    K.p(doc, "A credential that a graduate cannot easily hold, or that an employer cannot check "
             "without training, fails regardless of its cryptography. Rustemi et al. list staff "
             "training, cost of maintenance, diploma generation and the migration of existing "
             "certificates among the practical obstacles [1]. Open Badges 3.0 emphasises machine "
             "verifiability and portability for the same reason [12]. This project therefore treats "
             "verification as a public web page that requires no account, no wallet and no "
             "installation, and treats the registrar interface as the primary product rather than an "
             "afterthought.")

    K.h2(doc, "2.4 Prior systems and standards")
    K.p(doc, "**Blockcerts** is the most mature open ecosystem for blockchain-anchored credentials, "
             "originally developed at the MIT Media Lab with Learning Machine. It combines signed "
             "credential documents with blockchain anchoring and a recipient-controlled wallet model, "
             "and has moved towards alignment with W3C Verifiable Credentials, Data Integrity and "
             "DIDs [4][5]. Its enduring lesson is the separation of concerns: the blockchain "
             "provides verification evidence, while the credential itself lives with the recipient. "
             "Its difficulties — key management, issuer identification, revocation and "
             "interoperability — are the same difficulties this project confronts in reduced form.")
    K.p(doc, "**EduCTX** by Turkanović et al. generalises the problem from certificate authenticity "
             "to academic credit representation, using an Ark-based blockchain in which institutions "
             "are peers and ECTX tokens represent ECTS credits [7]. It demonstrates that shared "
             "academic records are feasible, and that the hard problems migrate to governance and "
             "identity once issuance is solved. **The Lifelong Learning Passport** of Gräther et al. "
             "extends the horizon to a learner's whole life, combining blockchain, smart contracts "
             "and storage so that credentials from many institutions can travel with the learner "
             "[8]; it makes clear that a credential ecosystem needs consistent data models and status "
             "handling to be more than a collection of isolated certificates.")
    K.p(doc, "**ZKBAR-V** is the most privacy-forward design reviewed. It combines two blockchains, "
             "IPFS, decentralised identifiers and zero-knowledge proofs so that verification can "
             "establish claims without disclosing the underlying record [6]. It is technically "
             "impressive and computationally expensive, and it illustrates the trade-off this project "
             "declines to take: selective disclosure is valuable, but it multiplies the components "
             "that must be correct before the system can be trusted.")
    K.p(doc, "Recent **Ethereum + IPFS** implementations reported in 2025 follow the pattern this "
             "project adopts: store the document off-chain, hash it, and anchor the hash, keeping the "
             "cost of the ledger proportional to the number of credentials rather than their size "
             "[10][11]. Their shared weakness is also instructive: anchoring a hash does not by "
             "itself solve identity, revocation, privacy, interoperability or the availability of the "
             "off-chain object.")
    K.p(doc, "Two **standards** complete the picture. W3C Verifiable Credentials 2.0 defines a data "
             "model, a proof mechanism and a status mechanism for machine-verifiable credentials "
             "[2]; it is not a blockchain specification, and a system may implement it with or "
             "without a ledger. Open Badges 3.0 profiles that model for achievement credentials and "
             "is widely used in education [12]. This project borrows their vocabulary — issuer, "
             "holder, verifier, credential subject, proof, status — while implementing a deliberately "
             "simpler signature scheme, and it does not claim conformance.")

    K.table(doc,
            ["System / standard", "Core architecture", "Main contribution", "Principal limitation"],
            [["Blockcerts [4][5]", "Signed credential documents plus blockchain anchoring",
              "Recipient-controlled, independently verifiable credentials",
              "Key management, revocation and adoption requirements"],
             ["EduCTX [7]", "Ark-based permissioned blockchain with institutional peers",
              "Shared representation of ECTS credits across institutions",
              "Credits rather than full credential lifecycle and privacy"],
             ["Lifelong Learning Passport [8]", "Blockchain, smart contracts and storage",
              "Portable record of achievements across a lifetime",
              "Needs ecosystem-wide standards and governance"],
             ["ZKBAR-V [6]", "Dual blockchain, IPFS, DIDs and zero-knowledge proofs",
              "Privacy-preserving verification with minimal disclosure",
              "High cryptographic and operational complexity"],
             ["Ethereum + IPFS studies [10][11]", "Smart contracts with IPFS-hosted documents",
              "Low on-chain footprint with tamper-evident verification",
              "Identity, status and governance remain out of scope"],
             ["W3C VC 2.0 [2]", "Standard data, proof and status model",
              "Interoperable, machine-verifiable credentials",
              "Not a ledger and not an implementation"],
             ["Open Badges 3.0 [12]", "VC-conformant achievement credential profile",
              "Standardised educational achievement claims",
              "Requires issuer adoption and identity infrastructure"]],
            "Comparison of prior systems and standards, synthesised from the cited work.",
            widths=[1.35, 1.75, 1.85, 1.6], font_size=8.8)

    K.h2(doc, "2.5 From research gaps to design requirements")
    K.p(doc, "The gaps identified in Section 2.3 were converted into the requirements that shaped the "
             "architecture. The mapping is recorded explicitly so that the design can be judged "
             "against the literature rather than against convenience.")
    K.table(doc,
            ["Gap in the literature", "Design response in this project", "Where it is realised"],
            [             ["Immutability mistaken for correctness",
              "Anchor integrity only; keep correctness inside institutional process and expose the issuer explicitly",
              "Sections 4.8, 5.7; issuer signature check"],
             ["Personal data on a public ledger",
              "Store transcripts off-chain; anchor digest, storage reference and status only",
              "Sections 4.4, 5.6; `academic_records` schema"],
             ["Credential lifecycle and revocation",
              "First-class Active / Revoked / Superseded states with a replacement link and atomic supersession",
              "Sections 4.5, 5.2; `supersedeRecord`"],
             ["Interoperability",
              "VC-shaped credential JSON with a documented canonicalisation rule, leaving a path to full VC conformance",
              "Sections 4.6, 5.4"],
             ["Cost and scale of anchoring",
              "One write per issuance or lifecycle change; all verification through a zero-gas view call",
              "Sections 4.5, 6.9"],
             ["Usability and adoption",
              "Public verification page requiring no wallet or account; QR and link sharing; registrar-centric interface",
              "Sections 5.9, 6.7"],
             ["Failure semantics and trust in reporting",
              "Explicit distinction between not found, tampered, revoked and unavailable, tested by an outage suite",
              "Sections 4.9, 6.6"]],
            "Traceability from reviewed gaps to implemented design responses.",
            widths=[1.85, 2.55, 2.1], font_size=8.8)

    K.h2(doc, "2.6 Summary")
    K.p(doc, "The literature supports a hybrid, standards-aware architecture: credentials stored "
             "off-chain under access control, a cryptographic commitment anchored on a ledger, an "
             "explicit status model, and a verification procedure that any third party can perform "
             "without the issuer's cooperation. It also warns, repeatedly, that blockchain alone "
             "solves none of identity, interoperability, governance or usability. The system "
             "described in the following chapters is a deliberate implementation of the defensible "
             "core of that position, with the remaining problems named rather than hidden.")


# ---------------------------------------------------------------------- chapter 3
def chapter3(doc):
    K.h1(doc, "3. System Analysis and Requirements")

    K.h2(doc, "3.1 Analysis of the existing process")
    K.p(doc, "The existing process at most institutions is a chain of custody rather than a system. "
             "Grades accumulate in departmental spreadsheets and are consolidated into a student "
             "information system; a transcript is generated on request, printed or exported as a "
             "PDF, stamped and signed; and a copy is filed. Verification is a separate workflow "
             "triggered by an external request, handled by staff who look up the student record and "
             "reply in writing.")
    K.p(doc, "Modelling that workflow exposes four failure surfaces. The artefact leaves the "
             "controlled environment at the moment of printing, so the institution no longer knows "
             "what version a verifier holds. The verification response is only as trustworthy as the "
             "channel that carried it. Corrections require locating and withdrawing every copy "
             "already issued, which in practice does not happen. And because the verification "
             "workflow is staffed, its latency and availability are tied to office hours. None of "
             "these are cryptographic problems, but each of them is a problem that a correctly "
             "designed anchoring layer can make observable.")
    K.p(doc, "It is equally important to record what must not change. Registrars need to correct "
             "grades and re-issue transcripts; students change names; programmes are revised; and "
             "access to a transcript must remain controlled. A design that makes records permanently "
             "unmodifiable would be unusable. The architecture therefore separates mutable academic "
             "data from immutable evidence, and models change as a new record linked to its "
             "predecessor rather than as an edit.")

    K.h2(doc, "3.2 The proposed system")
    K.p(doc, "The proposed system is a student information system in which academic record "
             "management is the product and verification is one integrated feature. It has three "
             "layers. A React single-page application provides dashboards for administrators, "
             "issuers and students, and a public verification page. A Node.js/Express API owns the "
             "business logic: it authenticates wallets, compiles transcripts, computes digests, "
             "encrypts and stores documents, coordinates transactions and cross-checks results. A "
             "MySQL 8 database holds the mutable academic data, while an Ethereum-compatible smart "
             "contract holds the immutable evidence.")
    K.p(doc, "The design rule that governs every decision is stated once and applied everywhere: "
             "**MySQL stores the mutable academic data; the chain stores only the immutable evidence — "
             "the SHA-256 digest, the storage reference, the issuer address, the issuance time and the "
             "lifecycle status.** No name, grade, programme or decryption key is written on-chain. "
             "Verification is a `view` function, so external verifiers pay nothing and never connect "
             "a wallet.")
    K.p(doc, "Four further principles follow from the literature. Issuance is *transactional*: a "
             "record is persisted as DRAFT, anchored, and promoted to ACTIVE only after the receipt "
             "is confirmed, so a failed transaction cannot be presented as an issued credential. "
             "Verification is *corroborative*: the ledger answer is compared with the database and "
             "the two must agree on both status and digest. Failures are *differentiated*: unknown, "
             "tampered, revoked and unavailable are distinct outcomes. And authorisation is "
             "*duplicated*: the API enforces roles, and the contract enforces an issuer whitelist "
             "that the API cannot override.")

    K.h2(doc, "3.3 Actors and responsibilities")
    K.table(doc,
            ["Actor", "Responsibilities", "Application role", "Ledger rights"],
            [["Registrar / faculty issuer",
              "Maintain academic data, issue credentials, revoke or correct records",
              "ISSUER", "Whitelisted issuer: may issue, revoke and supersede"],
             ["Administrator",
              "Manage users, programmes and courses; whitelist issuer wallets; inspect activity",
              "ADMIN", "Contract admin: may manage the issuer whitelist and transfer admin"],
             ["Student / holder",
              "View and download own credentials, share links and QR codes, raise corrections offline",
              "STUDENT", "None (read-only through the public interface)"],
             ["Verifier / public",
              "Check a credential by identifier, link or document hash; no account required",
              "Anonymous (VERIFIER role exists for institutional accounts)",
              "None (read-only `eth_call`)"],
             ["System signer",
              "Anchors transactions when the browser wallet is not whitelisted; server-side fallback",
              "Service account", "Whitelisted issuer in the default local deployment"]],
            "Actors, responsibilities and privileges.",
            widths=[1.25, 2.6, 1.3, 1.45], font_size=8.8)

    K.h2(doc, "3.4 Functional requirements")
    K.p(doc, "Requirements were written to be testable, and each is exercised by at least one of "
             "the suites described in Chapter 6.")
    K.table(doc,
            ["ID", "Requirement", "Verification"],
            [["FR-01", "Users authenticate by signing a challenge with an EVM wallet; no password is stored",
              "API suite, authentication group"],
             ["FR-02", "An unknown but valid wallet is auto-provisioned as a STUDENT",
              "API suite"],
             ["FR-03", "Only ADMIN and ISSUER roles may issue, revoke or correct records",
              "API suite (403 checks), contract tests"],
             ["FR-04", "Administrators maintain programmes, courses, enrolments and grades",
              "API suite, manual UI checks"],
             ["FR-05", "GPA is computed from completed enrolments using weighted grade points",
              "API suite, transcript assertions"],
             ["FR-06", "A credential can be issued from live academic data or from an uploaded PDF",
              "API suite, tamper suite"],
             ["FR-07", "Credential identifiers are unique in both the database and the ledger",
              "Contract tests, API suite"],
             ["FR-08", "The digest is SHA-256 over the exact finalised bytes; canonical JSON is key-sorted",
              "API suite, tamper suite"],
             ["FR-09", "Stored documents are encrypted at rest with AES-256-GCM",
              "API suite (ciphertext on disk)"],
             ["FR-10", "A record is ACTIVE only after the anchoring transaction is confirmed",
              "API suite, contract tests"],
             ["FR-11", "A failed anchoring transaction leaves a DRAFT record and an audit entry, retryable",
              "Chain-outage and failure-path tests"],
             ["FR-12", "Revocation updates both the database and the ledger",
              "API suite, contract tests"],
             ["FR-13", "Correction publishes a new record and supersedes the old one in one transaction",
              "API suite, contract tests"],
             ["FR-14", "Any party can verify by credential identifier without authentication",
              "API suite (public lookup)"],
             ["FR-15", "Any party can verify by document hash, detecting any modification",
              "Tamper suite (23 mutations)"],
             ["FR-16", "Verification reports database status, ledger status and a consistency flag",
              "API suite, tamper suite"],
             ["FR-17", "Students see and download only their own credentials",
              "API suite (scoping assertions)"],
             ["FR-18", "The verification page is reachable from a QR code or shareable link",
              "Manual UI checks"]],
            "Functional requirements and the evidence that each is met.",
            widths=[0.55, 4.0, 1.95], font_size=8.6, align_center_cols=(0,))

    K.h2(doc, "3.5 Non-functional requirements")
    K.table(doc,
            ["Attribute", "Requirement", "Outcome in the prototype"],
            [             ["Performance", "Verification must complete in well under a second on a local node",
              "Mean 54.7 ms, maximum 69 ms over ten trials (Section 6.9)"],
             ["Cost", "Verification must not consume gas or require a wallet",
              "Read-only `verifyRecord` view function; gas cost 0"],
             ["Reliability", "A ledger outage must not be reported as tampering",
              "Chain-outage suite passes 7/7 assertions (Section 6.6)"],
             ["Integrity", "Any modification of an issued document must be detectable",
              "23/23 mutations rejected; database tamper flagged (Section 6.5)"],
             ["Security", "Roles enforced in the application and independently in the contract",
              "API 401/403 tests; contract RBAC tests"],
             ["Confidentiality", "Documents protected at rest and in transit",
              "AES-256-GCM with per-record IV and authentication tag"],
             ["Auditability", "Every state change recorded with actor and transaction reference",
              "`audit_logs` table with eleven recorded action types"],
             ["Usability", "Verification requires no account, wallet or installation",
              "Public `/verify/:credId` page with QR and link sharing"],
             ["Portability", "Clear separation of standard and project-specific behaviour",
              "Environment-driven configuration; no secrets in the repository"],
             ["Maintainability", "Module boundaries aligned with responsibilities",
              "Contract, chain service, credential service, storage, routes"]],
            "Non-functional requirements and measured outcomes.",
            widths=[1.2, 2.75, 2.55], font_size=8.8)

    K.h2(doc, "3.6 Feasibility analysis")
    K.p(doc, "**Technical feasibility.** Every component used is mature and open source: Solidity "
             "0.8.24 with Hardhat for the contract, Node.js with Express and ethers v6 for the "
             "backend, React with Vite for the client, and MySQL 8 for storage. The only element "
             "that cannot be exercised without external credentials is IPFS pinning, which is "
             "therefore optional in the configuration.")
    K.p(doc, "**Economic feasibility.** On a local network the marginal cost of anchoring is zero. "
             "On a public network the design limits cost to one transaction per issuance and one per "
             "lifecycle change, with verification free. At the measured 101,548 gas per issuance, the "
             "cost of a cohort of a thousand credentials can be estimated directly from the network's "
             "gas price, and it scales linearly with the number of credentials rather than with their "
             "size.")
    K.p(doc, "**Operational feasibility.** The system runs on a developer workstation with MySQL and "
             "a local node. A registrar needs only a browser and a wallet; a verifier needs only a "
             "browser. The single operational obligation that differs from a conventional system is "
             "custody of the signing key and the document encryption key, which is discussed as a "
             "risk in Section 4.8.")
    K.p(doc, "**Legal and ethical feasibility.** The design deliberately keeps personal data "
             "off-chain, which reduces the tension between immutability and data-protection "
             "obligations. It does not eliminate that tension: a digest anchored for a revoked "
             "credential remains on the ledger, and linkability of records issued to the same student "
             "is possible. Both points are recorded as open issues in Section 7.3 rather than "
             "presented as solved.")

    K.h2(doc, "3.7 Development and deployment environment")
    K.p(doc, "The midterm proposal listed PostgreSQL and a purely W3C-oriented credential format as "
             "planned technologies. During implementation these were changed, and the changes are "
             "recorded here rather than quietly absorbed: MySQL 8 was used because the schema, "
             "foreign keys and JSON column behaviour were already validated in the department's "
             "laboratory environment, and the credential is a W3C-VC-*shaped* JSON document with a "
             "project-specific signature scheme, because implementing a conformant Data Integrity "
             "proof suite was judged out of scope for the available time. No standards conformance is "
             "claimed anywhere in this report.")
    K.table(doc,
            ["Component", "Technology and version", "Role"],
            [["Front end", "React 18.3, Vite 5.4, React Router 6.28, qrcode 1.5",
              "Dashboards, student portal, public verification page"],
             ["Backend", "Node.js 26 (requires 18+), Express 4.21, ethers 6.13, mysql2 3.11",
              "REST API, transcript compilation, hashing, chain coordination"],
             ["Database", "MySQL 8.0, InnoDB, utf8mb4",
              "Users, programmes, courses, enrolments, records, audit log"],
             ["Smart contract", "Solidity 0.8.24, optimizer enabled (200 runs), Hardhat 2.26",
              "Issuance, lifecycle, RBAC, zero-gas verification"],
             ["Ledger", "Local Hardhat node on port 8546, chain ID 31338; Sepolia configured",
              "Immutable evidence and events"],
             ["Storage", "Local encrypted uploads; IPFS via Pinata optional",
              "Credential documents off-chain"],
             ["Cryptography", "SHA-256, secp256k1 wallet signatures, AES-256-GCM",
              "Digest, issuer signature, encryption at rest"],
             ["Tooling", "Hardhat test runner, Mocha/Chai, custom Node test scripts, Git",
              "Unit, integration, adversarial and performance testing"]],
            "Implemented technology stack.",
            widths=[1.15, 3.0, 2.35], font_size=8.8)

    K.h2(doc, "3.8 Assumptions and constraints")
    K.bullets(doc, [
        "The issuing institution is honest at the point of issuance; the threat model protects "
        "against modification and unauthorised issuance after that point, not against a registrar "
        "entering a wrong grade.",
        "A verifier can obtain the exact bytes of the credential. Verification by digest is only "
        "meaningful if the verifier hashes the same byte sequence the issuer hashed, which is why the "
        "hashing rule is documented and the canonicalisation procedure is shared.",
        "The EVM node is reachable from the API host. If it is not, the system degrades to "
        "database-only reporting and says so explicitly.",
        "Wallet custody is the user's responsibility. Loss of an issuer wallet is a revocation event "
        "handled by the administrator, not a recoverable password reset.",
        "The prototype targets a single institution. Multiple institutions would require independent "
        "administrators and a governance agreement, which the contract structure anticipates but does "
        "not implement.",
    ])
