"""Small helper layer over python-docx for a consistent academic report."""
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x12, 0x20, 0x3D)
STEEL = RGBColor(0x2F, 0x4D, 0x7D)
GREY = RGBColor(0x55, 0x5F, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BODY_FONT = "Calibri"
HEAD_FONT = "Cambria"

# numbering state, driven by the most recent top-level heading
_STATE = {"chapter": None, "counts": {"figure": 0, "table": 0}}

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]

TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
    "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
]


def _insert_ordered(parent, element, order):
    """Insert element into parent at the schema-correct position."""
    tag = element.tag.split("}")[-1]
    idx = order.index(tag)
    for child in parent:
        child_tag = child.tag.split("}")[-1]
        if child_tag in order and order.index(child_tag) > idx:
            child.addprevious(element)
            return element
    parent.append(element)
    return element


def _p_border(paragraph, sides):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side, spec in sides.items():
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:space"), str(spec.get("space", 2)))
        el.set(qn("w:color"), spec.get("color", "C3CFE3"))
        borders.append(el)
    _insert_ordered(pPr, borders, PPR_ORDER)
    return borders


def _reset_numbering():
    _STATE["counts"]["figure"] = 0
    _STATE["counts"]["table"] = 0


def _label(kind):
    _STATE["counts"][kind] += 1
    n = _STATE["counts"][kind]
    chapter = _STATE["chapter"]
    return f"{chapter}.{n}" if chapter else str(n)


# --------------------------------------------------------------------- fields
def _field(paragraph, instr, placeholder="", bold=False, size=None, color=None):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run._r.append(instr_el)

    run = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    run = paragraph.add_run(placeholder)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color

    run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def enable_update_fields(document):
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    # CT_Settings requires updateFields before compat / rsids blocks.
    for child in settings:
        tag = child.tag.split("}")[-1]
        if tag in ("compat", "docVars", "rsids", "mathPr", "themeFontLang",
                   "clrSchemeMapping", "shapeDefaults", "decimalSymbol"):
            child.addprevious(update)
            return
    settings.append(update)


# --------------------------------------------------------------------- styles
def _set_font(style, name, size=None, bold=None, color=None, italic=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def _style_heading(style, size, color=NAVY, before=14, after=6):
    _set_font(style, HEAD_FONT, size=size, bold=True, color=color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.line_spacing = 1.15


def build_document():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(0.9)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    _set_font(normal, BODY_FONT, size=11, color=RGBColor(0x1A, 0x1A, 0x1A))
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _style_heading(doc.styles["Heading 1"], 17, before=0, after=12)
    _style_heading(doc.styles["Heading 2"], 13.5, color=STEEL)
    _style_heading(doc.styles["Heading 3"], 11.5, color=STEEL, before=10, after=4)

    cap = doc.styles["Caption"]
    _set_font(cap, BODY_FONT, size=9, italic=True, color=GREY)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in ("FigCaption", "TabCaption"):
        if name not in [s.name for s in doc.styles]:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles["Caption"]
            style.next_paragraph_style = doc.styles["Normal"]

    # Front-matter headings look like chapter headings but stay out of the TOC.
    if "FrontHeading" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("FrontHeading", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 1"]
        style.next_paragraph_style = doc.styles["Normal"]
        pPr = style.element.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "9")
        _insert_ordered(pPr, outline, PPR_ORDER)

    return doc


# ------------------------------------------------------------- rich text run
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_rich_text(paragraph, text, size=None, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 1.0)
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


# ------------------------------------------------------------------ headings
def h1(doc, text, page_break=True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    _p_border(p, {"bottom": {"sz": 8, "space": 4, "color": "2F4D7D"}})

    m = re.match(r"^(\d+)\.\s", text)
    if m:
        _STATE["chapter"] = m.group(1)
    elif text.startswith("Appendix "):
        _STATE["chapter"] = text[len("Appendix "):].strip()[0]
    else:
        _STATE["chapter"] = None
    _reset_numbering()
    return p


def front_heading(doc, text, page_break=True):
    if page_break:
        doc.add_page_break()
    para = doc.add_paragraph(text, style="FrontHeading")
    para.paragraph_format.space_before = Pt(0)
    _p_border(para, {"bottom": {"sz": 8, "space": 4, "color": "2F4D7D"}})
    return para


def h2(doc, text):
    return doc.add_paragraph(text, style="Heading 2")


def h3(doc, text):
    return doc.add_paragraph(text, style="Heading 3")


# ------------------------------------------------------------------ body text
def p(doc, text, size=None, align=None, space_after=None, italic=False, color=None):
    para = doc.add_paragraph()
    add_rich_text(para, text, size=size, color=color)
    if italic:
        for run in para.runs:
            run.italic = True
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if space_after is not None:
        para.paragraph_format.space_after = Pt(space_after)
    return para


def bullets(doc, items, numbered=False, size=None):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        para = doc.add_paragraph(style=style)
        add_rich_text(para, item, size=size)
        para.paragraph_format.line_spacing = 1.35
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.left_indent = Inches(0.3)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def code_block(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(10)
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.keep_together = True
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _p_border(para, {
        "left": {"sz": 18, "space": 6, "color": "2F4D7D"},
        "top": {"sz": 4, "space": 6},
        "bottom": {"sz": 4, "space": 6},
        "right": {"sz": 4, "space": 6},
    })
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F6FB")
    _insert_ordered(pPr, shd, PPR_ORDER)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if i < len(lines) - 1:
            run.add_break()
    return para


# ------------------------------------------------------------------- captions
def _caption(doc, kind, prefix, number, text):
    style = "FigCaption" if kind == "figure" else "TabCaption"
    para = doc.add_paragraph(style=style)
    run = para.add_run(f"{prefix} {number}.")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = NAVY
    run.font.name = BODY_FONT
    run = para.add_run(f" {text}")
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = GREY
    run.font.name = BODY_FONT
    return para


def figure(doc, path, caption, width=6.2):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(path, width=Inches(width))
    _caption(doc, "figure", "Figure", _label("figure"), caption)


# --------------------------------------------------------------------- tables
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
              "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    _insert_ordered(tcPr, shd, TCPR_ORDER)


def _cell_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "C3CFE3")
        borders.append(el)
    _insert_ordered(tblPr, borders, TBLPR_ORDER)


def _cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insert_ordered(tblPr, mar, TBLPR_ORDER)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def table(doc, headers, rows, caption, widths=None, font_size=9.2, align_center_cols=()):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.keep_with_next = True

    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _cell_borders(tbl)
    _cell_margins(tbl)

    hdr = tbl.rows[0]
    _repeat_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.line_spacing = 1.05
        para.paragraph_format.space_after = Pt(0)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = WHITE
        run.font.name = BODY_FONT
        _shade(cell, "2F4D7D")

    for r_idx, row in enumerate(rows):
        cells = tbl.add_row().cells
        _no_split(tbl.rows[-1])
        for c_idx, text in enumerate(row):
            cell = cells[c_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing = 1.08
            para.paragraph_format.space_after = Pt(0)
            if c_idx in align_center_cols:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_rich_text(para, str(text), size=font_size)
            for run in para.runs:
                if run.font.name != "Consolas":
                    run.font.name = BODY_FONT
            if r_idx % 2 == 1:
                _shade(cell, "F3F6FB")

    if widths:
        total = sum(widths)
        usable = 6.4
        if total > usable:
            widths = [w * usable / total for w in widths]
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)

    _caption(doc, "table", "Table", _label("table"), caption)
    return tbl


# --------------------------------------------------------- header and footer
def add_header_footer(doc, header_text):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    para = header.paragraphs[0]
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run(header_text)
    run.font.size = Pt(7.5)
    run.font.color.rgb = GREY
    run.font.name = BODY_FONT
    _p_border(para, {"bottom": {"sz": 4, "space": 2}})

    footer = section.footer
    para = footer.paragraphs[0]
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("Page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    _field(para, "PAGE", "1", size=8.5, color=GREY)
    run = para.add_run(" of ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    _field(para, "NUMPAGES", "1", size=8.5, color=GREY)


def toc(doc, title, instr, placeholder="Open this document in Word and update this field to populate the list."):
    para = doc.add_paragraph(title, style="FrontHeading")
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(10)
    _p_border(para, {"bottom": {"sz": 8, "space": 4, "color": "2F4D7D"}})

    body = doc.add_paragraph()
    body.paragraph_format.line_spacing = 1.3
    _field(body, instr, placeholder, size=10)


def blank(doc, count=1):
    for _ in range(count):
        doc.add_paragraph()


def centered(doc, text, size=11, bold=False, italic=False, color=None, space_after=4,
             font=BODY_FONT):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.15
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    if color:
        run.font.color.rgb = color
    return para
